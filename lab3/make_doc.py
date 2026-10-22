from docx import Document
from docx.shared import Pt, Inches
H = "We chose Microservices Architecture for the Telemedicine Slot Booking & Prescription Portal."
items = [
("Architectural Choice",
 "Microservices. The Lab 3 design covers Set 1 (Doctor Discovery & Slot Booking) and its hand-offs: a Patient Web/Mobile App, an API Gateway, independently deployable Doctor Search, Booking, Payment and Notification services, each Booking/Search service with its own database, and the external Payment Gateway. They communicate only through the interfaces ISearchDoctors, IBooking, IAvailability, IPayment, INotify, ISlotStore, IDoctorData and IGateway."),
("Reason 1 - Independent scaling (NFR-002, NFR-003)",
 "Search is read-heavy and bursty, while Booking is write-heavy and concurrency-critical (FR-003: a booked slot must vanish for all other patients within 2 seconds and never be double-booked). Separate services scale on their own, and the Booking Service can enforce slot holds and reservation atomically in its own database. The portal must also support 500 concurrent video sessions at under 300 ms latency, which a single layered or client-server deployment would not scale to cheaply."),
("Reason 2 - Fault isolation and replaceability",
 "The portal targets 99.9% uptime with no confirmed booking lost. If the Notification Service or the external Payment Gateway fails, search and slot-hold still work and the failed step can be retried (the UC-02 payment-declined flow releases the hold after 2 attempts). Payment and notification can also be swapped for another provider without changing Search or Booking."),
("Security Advantage",
 "All patient traffic enters through the API Gateway, which authenticates users and terminates TLS 1.3 (NFR-001); internal services are never exposed publicly. Card data is confined to the Payment Service and the external gateway, and each database is reachable only through its owning service, limiting the blast radius of a breach of patient health data."),
("Performance Benefit",
 "Each service is tuned independently: Doctor Search can cache specialty/availability results and use read replicas, so the 7-day availability filter (FR-002) responds quickly, while Booking keeps a small, fast transactional store for slot holds."),
("Styles considered",
 "Layered: a single deployable unit cannot scale search, booking and video separately, and one failure affects every function. Client-Server: a central server is a single point of failure and a bottleneck, conflicting with NFR-003. The trade-off accepted is higher operational complexity and eventual consistency between services, mitigated by the gateway and a message queue for notifications."),
]
d = Document()
for s in d.sections:
    s.top_margin = s.bottom_margin = Inches(0.7); s.left_margin = s.right_margin = Inches(0.85)
d.styles['Normal'].font.name = 'Calibri'; d.styles['Normal'].font.size = Pt(10.5)
d.add_heading("Lab 3: Architecture Selection Justification", 1)
p = d.add_paragraph(); p.add_run(H).bold = True
for t, b in items:
    p = d.add_paragraph(); p.add_run(t + ": ").bold = True; p.add_run(b)
    p.paragraph_format.space_after = Pt(5)
d.save("Lab3_Telemedicine_Justification.docx")
